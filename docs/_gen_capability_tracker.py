# -*- coding: utf-8 -*-
# 警告：重跑会覆盖 docs/系统能力跟踪表.xlsx 里已填写的进度，仅在能力清单本身变更时使用。
# Generate docs/系统能力跟踪表.xlsx — tracker for WMS capabilities.
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from datetime import date

OUT = r"e:\code\__self\v4-wms\docs\系统能力跟踪表.xlsx"

# —— 状态口径 ——
# 原型设计 / UI设计：0 期 HTML 原型与 UI设计.md 已覆盖的标「已完成」
# 前端 / 后端 / 对接 / 测试：正式 React + Java 工程尚未开工，标「未开始」
# 纯工程/信创项：原型与 UI 标「不适用」
DONE, TODO, NA = "已完成", "未开始", "不适用"

PRIMARY = "1B4F8A"
PRIMARY_DARK = "153D6B"
ACCENT = "2F80C4"
SUCCESS = "2E7D4F"
WARNING = "C47A12"
DANGER = "B42318"
SURFACE = "FFFFFF"
BG = "F0F3F7"
BORDER_C = "D8DEE8"
TEXT = "1A2332"
MUTED = "5A6A7E"

fill_primary = PatternFill("solid", fgColor=PRIMARY)
fill_primary_dark = PatternFill("solid", fgColor=PRIMARY_DARK)
fill_head = PatternFill("solid", fgColor="1B4F8A")
fill_sub = PatternFill("solid", fgColor="E8EEF6")
fill_white = PatternFill("solid", fgColor=SURFACE)
fill_bg = PatternFill("solid", fgColor=BG)
fill_done = PatternFill("solid", fgColor="E3F2E9")
fill_doing = PatternFill("solid", fgColor="FFF4D6")
fill_todo = PatternFill("solid", fgColor="F3F5F8")
fill_block = PatternFill("solid", fgColor="FDECEA")
fill_na = PatternFill("solid", fgColor="EEF3F9")
fill_wait = PatternFill("solid", fgColor="E7F1FB")
fill_p0 = PatternFill("solid", fgColor="FDECEA")
fill_p1 = PatternFill("solid", fgColor="FFF4D6")
fill_p2 = PatternFill("solid", fgColor="EEF3F9")
fill_title_bar = PatternFill("solid", fgColor="0D2B4A")

font_title = Font(name="微软雅黑", size=18, bold=True, color="FFFFFF")
font_h = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
font_h2 = Font(name="微软雅黑", size=12, bold=True, color=PRIMARY)
font_body = Font(name="微软雅黑", size=10, color=TEXT)
font_bold = Font(name="微软雅黑", size=10, bold=True, color=TEXT)
font_muted = Font(name="微软雅黑", size=9, color=MUTED)
font_white = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
font_kpi = Font(name="微软雅黑", size=20, bold=True, color=PRIMARY)

thin = Border(
    left=Side(style="thin", color=BORDER_C),
    right=Side(style="thin", color=BORDER_C),
    top=Side(style="thin", color=BORDER_C),
    bottom=Side(style="thin", color=BORDER_C),
)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center", wrap_text=True)

STATUSES = ["未开始", "进行中", "已完成", "阻塞", "不适用"]
STAGES = ["未开始", "原型中", "UI设计中", "待开发", "开发中", "对接中", "测试中", "已完成", "阻塞", "不适用"]

HEADERS = [
    "编号", "交付批次", "目标节点", "模块", "页面", "能力名称", "能力说明 / 验收要点",
    "关联接口", "优先级",
    "原型设计", "UI设计", "前端开发", "后端开发", "前后端对接", "测试",
    "阶段", "前端负责人", "后端负责人", "测试负责人", "计划完成", "实际完成", "阻塞原因", "备注",
]


def C(batch, node, module, page, name, desc, api, pri="P0", proto=DONE, ui=DONE, note="", fe=TODO, be=TODO, integ=TODO, test=TODO):
    return (batch, node, module, page, name, desc, api, pri, proto, ui, fe, be, integ, test, note)


CAPS = [
    # —— 公共框架 / 启动 ——
    C("启动", "D5", "公共框架", "工程", "需求基线确认", "输出需求确认单、流程说明、数据字典、权限矩阵、接口清单并双方签字", "—", "P0", NA, NA, "合同签订后 5 日内"),
    C("启动", "D5", "公共框架", "工程", "信创环境规划", "确认 CPU 架构、JDK 版本、Vastbase 实例授权与运行基线", "—", "P0", NA, NA),
    C("启动", "D5", "公共框架", "工程", "工程骨架可运行", "Java Spring Boot + Flyway（PG 方言）+ JWT；前端登录鉴权、布局与权限路由", "POST /api/auth/login", "P0", NA, NA),
    C("启动", "D5", "公共框架", "登录", "多角色登录", "仓管员 / 关务员 / 管理员登录；按注册角色进入对应默认页", "POST /api/auth/login", "P0"),
    C("启动", "D5", "公共框架", "登录", "登录态保持与退出", "session 保持；退出清除令牌；角色切换刷新菜单", "POST /api/auth/logout", "P0"),
    C("启动", "D5", "公共框架", "全局布局", "三栏布局与导航", "左侧导航 + 顶栏 + 内容区；侧栏分组：概览 / 基础数据 / 仓储作业 / 生产与关务 / 系统", "—", "P0"),
    C("启动", "D5", "公共框架", "全局布局", "基础组件库", "Button / Card / Input / Select / Modal / Toast / Badge / Table / Tabs / Upload / Empty / Loader", "—", "P0"),
    C("启动", "D5", "公共框架", "全局布局", "统一响应与错误提示", "接口统一 {success,data,message}；校验失败、权限不足、网络错误可感知", "—", "P0", NA, NA),
    C("第一批", "D10", "公共框架", "权限", "角色权限矩阵", "仓管员：入出库/移库/盘点/库存查询；关务员菜单默认全开；管理员全量含档案与设置", "GET/PUT /api/roles/{id}/permissions", "P0"),
    C("第一批", "D10", "公共框架", "权限", "越权路由拦截", "前端藏菜单 + 后端接口级鉴权；无权限跳转默认页", "—", "P0"),
    C("第一批", "D10", "公共框架", "权限", "海关专用账号审计标识", "关务员操作在审计日志中单独标识", "GET /api/audit-logs", "P0"),
    C("第一批", "D10", "公共框架", "等保基线", "密码策略", "强密码规则（长度/复杂度），前后端双重校验", "POST /api/auth/login", "P0", NA, NA),
    C("第一批", "D10", "公共框架", "等保基线", "登录失败锁定", "连续失败锁定账号/IP，可配置阈值与时长", "POST /api/auth/login", "P0", NA, NA),
    C("第一批", "D10", "公共框架", "等保基线", "会话超时", "空闲超时自动退出并提示重新登录", "—", "P0"),
    C("第一批", "D10", "公共框架", "工程", "信创最小冒烟", "UOS Server + Vastbase 部署 Java 进程，完成登录 + 一条档案维护", "GET /api/health", "P0", NA, NA, "开发便利环境与验收环境双轨"),

    # —— 仓库 ——
    C("第一批", "D10", "仓库管理", "仓库管理", "仓库列表", "展示 1/2/4/5/6 仓及码头仓库；编号、库容、面积㎡、堆位数、占用率、状态", "GET /api/warehouses", "P0"),
    C("第一批", "D10", "仓库管理", "仓库管理", "仓库新建/编辑", "维护总库容与面积㎡；面积仅档案展示可编辑，不参与占用计算", "POST /api/warehouses", "P0"),
    C("第一批", "D10", "仓库管理", "仓库管理", "仓库启停", "停用仓库不参与作业分配，平面图可标停用仓", "POST /api/warehouses", "P1"),

    # —— 堆位 ——
    C("第一批", "D10", "堆位管理", "堆位管理", "堆位列表（单仓）", "仅展示当前所选仓库（默认 1 仓）；编码、仓库、片区标识、网格、库容、状态", "GET /api/warehouses/stacks", "P0"),
    C("第一批", "D10", "堆位管理", "堆位管理", "堆位新建/编辑", "编码由仓库编号+堆位名称自动拼接（如 1#A1）只读；不展示网格字段", "POST /api/warehouses/stacks", "P0"),
    C("第一批", "D10", "堆位管理", "堆位管理", "档案字段收敛", "不展示报关单号、生产批次、占用率、查验状态", "GET /api/warehouses/stacks", "P0"),
    C("第一批", "D10", "堆位管理", "堆位管理", "网格取自平面图", "rows/cols 由园区平面图该堆位占格计算；列表只读展示", "GET /api/floor-plans", "P0"),
    C("第一批", "D10", "堆位管理", "堆位管理", "片区标识三类", "仅原料区 / 待检区 / 成品区，随在库批次查验自动变更", "PUT /api/warehouses/stacks/{id}/area-label", "P0"),
    C("第一批", "D10", "堆位管理", "堆位管理", "查验完成改成品区", "待检区堆位执行查验完成后改为成品区", "PUT /api/warehouses/stacks/{id}/area-label", "P0"),
    C("第一批", "D10", "堆位管理", "堆位管理", "片区变更留痕", "片区标识变更写入操作日志（改前改后）", "GET /api/audit-logs", "P0"),

    # —— 物料 ——
    C("第一批", "D10", "物料档案", "物料档案", "物料列表", "编码、名称、货物类型、类型、矿源", "GET /api/materials", "P0"),
    C("第一批", "D10", "物料档案", "物料档案", "新建/编辑仅基础信息", "仅编码、名称、货物类型、类型；报备矿填矿源。不含干湿重规则、默认账册、单证样例、原产国、品质参数、附件", "POST /api/materials", "P0"),
    C("第一批", "D10", "物料档案", "物料档案", "货物类型规则", "原料：达标矿/报备矿；成品一律混成品", "POST /api/materials", "P0"),
    C("第一批", "D10", "物料档案", "物料档案", "物料导出 Excel", "导出编码/名称/货物类型/类型/矿源", "GET /api/materials/export", "P1"),

    # —— 矿源备案 ——
    C("第一批", "D10", "矿源备案", "物料档案 / 矿源备案", "备案 CRUD 与附件", "矿山、原产国、备案数量、核销/剩余/用量%、附件；一矿源多期备案", "GET/POST /api/material-sources/filings", "P0"),
    C("第一批", "D10", "矿源备案", "物料档案 / 矿源备案", "备案列表筛选", "按矿源、编号、状态、备案日期起止查询", "GET /api/material-sources/filings", "P0"),
    C("第一批", "D10", "矿源备案", "物料档案 / 矿源备案", "点击编号查看票货", "该备案下全部票货：报关单号、生产批次号、委托方、重量、堆位、状态", "GET /api/material-sources/filings/{code}/lots", "P0", DONE, DONE, "待确认：详情「已核销」指生产消耗还是成品完整出库"),
    C("第一批", "D10", "矿源备案", "物料档案 / 矿源备案", "备案列表导出", "矿源备案列表导出 Excel", "GET /api/material-sources/filings/export", "P1"),
    C("第一批", "D10", "矿源备案", "物料档案 / 矿源备案", "票货导出", "备案下全部票货导出 Excel", "GET /api/material-sources/filings/{code}/lots/export", "P1"),
    C("第一批", "D10", "矿源备案", "物料档案 / 矿源备案", "入库自动匹配备案号", "报备矿入库按矿源自动匹配当前有效且有剩余的备案号", "POST /api/inbound", "P0"),
    C("第一批", "D10", "矿源备案", "物料档案 / 矿源备案", "干重扣减剩余量", "报备矿入库按干重扣减备案剩余（剩余−本次干重）", "POST /api/inbound", "P0"),
    C("第一批", "D10", "矿源备案", "物料档案 / 矿源备案", "用量阈值预警", "用量达系统参数阈值（默认 80%）触发预警", "GET /api/inventory/alerts", "P0"),
    C("第一批", "D10", "矿源备案", "物料档案 / 矿源备案", "关联生产批次号", "备案详情展示关联生产批次号", "GET /api/material-sources/filings", "P1"),
    C("第二批", "D20", "矿源备案", "库存查询 / 本票详情", "本票成品去向", "原料本票详情增加成品页：产量、库存、出库", "GET /api/inventory/batches/{batchNo}", "P1"),

    # —— 往来主体 ——
    C("第一批", "D10", "往来主体", "往来主体", "主体 CRUD", "委托方、物流账册主体手工维护；编码、名称、角色、电话、状态、来源；不含统一社会信用代码", "GET/POST /api/partners", "P0"),
    C("第一批", "D10", "往来主体", "往来主体", "委托方下拉", "五矿有色 / 广西金川 / 广西南国，贯穿出入库/移库/库存/生产", "GET /api/partners/consignors", "P0"),
    C("第二批", "D20", "往来主体", "出库管理", "流向企业自动建档", "出库报关单「流向企业」字段自动写入往来主体，无需手工新建", "POST /api/outbound", "P0"),

    # —— 保税账册 ——
    C("第一批", "D10", "保税账册", "保税账册", "账册 CRUD", "编号、名称、经营单位、货值、报关单号、状态", "GET/POST /api/bonded-books", "P0"),
    C("第一批", "D10", "保税账册", "保税账册", "新旧账衔接", "新账须填旧账预警日期；旧账只出不进", "POST /api/bonded-books", "P0"),
    C("第一批", "D10", "保税账册", "保税账册", "到期票转入新账", "超过预警日后未出库票的入库数据转入新账", "POST /api/bonded-books", "P0"),
    C("第一批", "D10", "保税账册", "保税账册", "账册列表筛选", "关键字 + 旧账预警起止日期", "GET /api/bonded-books", "P1"),
    C("第一批", "D10", "保税账册", "保税账册", "点击编号查看票矿", "账册下全部出入库票矿", "GET /api/bonded-books/{code}/lots", "P0"),
    C("第一批", "D10", "保税账册", "保税账册", "票矿导出 Excel", "账册票矿导出", "GET /api/bonded-books/{code}/lots/export", "P1"),

    # —— 系统设置 ——
    C("第一批", "D10", "系统设置", "系统设置 / 账号管理", "全站账号注册", "新建管理员/仓管员/关务员；列表无所属关区", "GET/POST /api/auth/users", "P0"),
    C("第一批", "D10", "系统设置", "系统设置 / 角色权限", "配置权限弹窗", "按菜单分组勾选可见模块，预勾角色默认权限，保存后侧栏同步", "GET/PUT /api/roles/{id}/permissions", "P0"),
    C("第一批", "D10", "系统设置", "系统设置 / 系统参数", "库容预警参数", "片区占用达额定库容 70% 触发，可配置", "待列入 API：系统参数 GET/PUT", "P0"),
    C("第一批", "D10", "系统设置", "系统设置 / 系统参数", "原料库龄参数", "原料在库超过 90 天触发，可配置", "待列入 API：系统参数 GET/PUT", "P0"),
    C("第一批", "D10", "系统设置", "系统设置 / 系统参数", "码头仓库预警参数", "码头仓库货品超期天数，默认 7 天", "待列入 API：系统参数 GET/PUT", "P0"),
    C("第一批", "D10", "系统设置", "系统设置 / 系统参数", "矿源备案用量预警参数", "用量百分比阈值，默认 80%", "待列入 API：系统参数 GET/PUT", "P1"),
    C("第一批", "D10", "系统设置", "系统设置 / 系统参数", "品质参数配置", "可新增/删减品质项；入库收货与 OCR 品质字段随配置同步；铜为必填", "待列入 API：系统参数 GET/PUT", "P0"),
    C("第一批", "D10", "系统设置", "系统设置 / 系统参数", "日志永久留存", "不提供年限配置；移除盘点差异自动生成等已废弃参数", "—", "P1"),

    # —— 平面图 ——
    C("第一批", "D10", "平面图设置", "系统设置 / 平面图设置", "园区画板与网格密度", "左画板右工具；网格密度为行×列，随保存写入 gridRows/gridCols", "GET/PUT /api/floor-plans", "P0", DONE, DONE, "0 期原型接口已通（Python），正式 Java 需重做"),
    C("第一批", "D10", "平面图设置", "系统设置 / 平面图设置", "缩放与平移", "滚轮 1x–8x；右上角 −/%/＋；空格/中键/空白处拖拽平移；平移不改宽高", "—", "P0"),
    C("第一批", "D10", "平面图设置", "系统设置 / 平面图设置", "绘制园区对象", "5 个启用仓 + 停用仓 + 办公楼 + 道路/灌木铺贴；仓内嵌套生产线/门/堆位", "PUT /api/floor-plans", "P0"),
    C("第一批", "D10", "平面图设置", "系统设置 / 平面图设置", "标注与默认填色", "仓号标在矩形下方居中；堆位名左下角不含仓号前缀、白底；仓库默认 #f8e4c5、门默认 #9ebce9；门置顶", "PUT /api/floor-plans", "P0"),
    C("第一批", "D10", "平面图设置", "系统设置 / 平面图设置", "编辑交互", "点选拖拽、四角缩放、网格吸附四边、Ctrl+Z/Y 撤销重做；松手贴边不重叠", "—", "P0"),
    C("第一批", "D10", "平面图设置", "系统设置 / 平面图设置", "生产区文字自适应", "生产线/功能区名称随块缩放，允许换行，窄条竖排", "—", "P1"),
    C("第一批", "D10", "平面图设置", "系统设置 / 平面图设置", "平面图保存读取", "园区键 park；短 JSON 回包；旧单仓接口兼容", "GET/PUT /api/floor-plans", "P0", DONE, DONE, "0 期原型接口已通（Python）"),

    # —— 入库 ——
    C("第二批", "D20", "入库管理", "入库管理", "入库列表与筛选", "时间范围、货运方式、矿种筛选；状态暂存码头/待收货/已入库/混成品；无片区标识列", "GET /api/inbound", "P0"),
    C("第二批", "D20", "入库管理", "入库管理", "列表业务字段", "矿种、装运方式（集装箱/散货）、集装箱柜数、金额及币种（¥/$）", "GET /api/inbound", "P0"),
    C("第二批", "D20", "入库管理", "入库管理", "新建预约收货", "船名、提单号、报关单、核注清单、金额、矿种、装运方式、柜数、委托方；预约与到货不含五项有害元素", "POST /api/inbound", "P0"),
    C("第二批", "D20", "入库管理", "入库管理", "预定堆位搜索", "预定堆位移到存放位置上方，支持编码/仓名关键字搜索", "GET /api/warehouses/stacks", "P0"),
    C("第二批", "D20", "入库管理", "入库管理", "圈选存放格子", "须选堆位并拖选/单击圈选存放格子，提交 stackCode + cells[{r,c}]", "POST /api/inbound", "P0"),
    C("第二批", "D20", "入库管理", "入库管理", "OCR 提单/报关单回填", "上传后回填船名、提单号、报关单号、核注清单号、金额币种、物料；可打开识别详情修正", "POST /api/ocr/recognize", "P0"),
    C("第二批", "D20", "入库管理", "入库管理", "到货收货登记", "到货确认，状态由暂存码头/待收货推进", "POST /api/inbound", "P0"),
    C("第二批", "D20", "入库管理", "入库管理", "称重干湿重", "录入湿重/干重/水分", "POST /api/inbound", "P0"),
    C("第二批", "D20", "入库管理", "入库管理", "OCR 重量单回填", "识别净湿重/干重/水分并回填", "POST /api/ocr/recognize", "P0"),
    C("第二批", "D20", "入库管理", "入库管理", "OCR 品质证书回填", "识别 Cu/Ag/Au/As/Pb/Cd/F/Hg，铜必填；字段随系统参数同步", "POST /api/ocr/recognize", "P0"),
    C("第二批", "D20", "入库管理", "入库管理", "保税账册匹配", "收货时匹配保税账册", "POST /api/inbound", "P0"),
    C("第二批", "D20", "入库管理", "入库管理", "上架写库存流水", "原料库存批次取报关单号；按片区分配堆位；禁止脱离单据改数", "POST /api/inbound", "P0"),
    C("第二批", "D20", "入库管理", "入库管理", "本票单证详情", "已入库票查看报关单、重量证书、品质证书及八项品质", "GET /api/inventory/batches/{batchNo}", "P0"),

    # —— 移库 ——
    C("第二批", "D20", "库内移库", "库内移库", "移库列表筛选", "按委托方、场景、起止日期筛选；列含船名、柜数、起止到时；无操作人/详情", "GET /api/transfers", "P0"),
    C("第二批", "D20", "库内移库", "库内移库", "先选堆位再选批次", "新建移库先选原堆位再选矿批次", "POST /api/transfers", "P0"),
    C("第二批", "D20", "库内移库", "库内移库", "批次信息只读带出", "委托方/报关单号/生产批次/物料/船名/柜数以表格只读展示", "POST /api/transfers", "P0"),
    C("第二批", "D20", "库内移库", "库内移库", "整票移位", "整票从原堆位迁到目标堆位", "POST /api/transfers", "P0"),
    C("第二批", "D20", "库内移库", "库内移库", "拆票移位", "按重量拆分移位，原票剩余与新票并存", "POST /api/transfers", "P0"),
    C("第二批", "D20", "库内移库", "库内移库", "批量移库跨仓", "可勾选全部仓库矿批次；按仓库/原堆位/委托方/关键字查询后提交；每票一条不合并", "POST /api/transfers/batch", "P0"),
    C("第二批", "D20", "库内移库", "库内移库", "移库时间精确到时", "起止到时，不含分", "POST /api/transfers", "P1"),
    C("第二批", "D20", "库内移库", "库内移库", "移库事务库存更新", "同一事务扣减原堆位、增加目标堆位；失败整体回滚；目标片区随在库批次状态更新", "POST /api/transfers", "P0"),

    # —— 库存 ——
    C("第二批", "D20", "库存查询", "库存查询", "多维查询", "矿源/账册/原料成品/堆位/片区标识/时间；列表不含类型列", "GET /api/inventory", "P0"),
    C("第二批", "D20", "库存查询", "库存查询", "时间范围到时", "本月，或自定义起止精确到时（如 2026/8/01/8时–2026/8/2/16时）", "GET /api/inventory", "P0"),
    C("第二批", "D20", "库存查询", "库存查询", "货物类型与空值规则", "原料达标矿/报备矿、成品混成品；成品原产国为 —；已出库堆位为 —", "GET /api/inventory", "P0"),
    C("第二批", "D20", "库存查询", "库存查询", "本票品质八项", "Cu/Ag/Au/As/Pb/Cd/F/Hg，品质属于上架后本票库存", "GET /api/inventory", "P0"),
    C("第二批", "D20", "库存查询", "库存查询", "原料本票详情", "报关单、重量证书、品质证书 + 成品去向（产量/库存/出库）", "GET /api/inventory/batches/{batchNo}", "P0"),
    C("第二批", "D20", "库存查询", "库存查询", "成品本票详情", "出库报关单、品质证书；未出库时出库报关单为空", "GET /api/inventory/batches/{batchNo}", "P0"),
    C("第二批", "D20", "库存查询", "库存查询", "干湿重重量台账", "按报关单号或生产批次号穿透业务流水；出库行含流向企业", "GET /api/inventory/batches/{batchNo}/ledger", "P0"),

    # —— 盘点 ——
    C("第二批", "D20", "盘点管理", "盘点管理", "盘点列表", "盘点单列表，无类型字段", "GET /api/stocktakes", "P0"),
    C("第二批", "D20", "盘点管理", "盘点管理", "新建盘点", "范围可多选仓库与堆位；可上传附件；不回写库存、不生成账实差异", "POST /api/stocktakes", "P0"),
    C("第二批", "D20", "盘点管理", "盘点管理", "月度在库导出", "按月份导出各委托方在库湿/干重、堆位、货物类型、提单号、报关单号（月末快照）", "GET /api/stocktakes/onhand-export", "P0"),

    # —— 预警 ——
    C("第二批", "D20", "库存预警", "库存预警", "库容预警", "片区占用达额定库容 70% 触发，可跳转业务页", "GET /api/inventory/alerts", "P0"),
    C("第二批", "D20", "库存预警", "库存预警", "原料库龄预警", "原料在库超过 90 天触发", "GET /api/inventory/alerts", "P0"),
    C("第二批", "D20", "库存预警", "库存预警", "码头仓库超期预警", "码头仓库货品超过参数天数（默认 7）触发", "GET /api/inventory/alerts", "P0"),
    C("第一批", "D10", "库存预警", "库存预警", "账册预警转入", "旧账预警日后未出库票转入新账的预警展示", "GET /api/inventory/alerts", "P1"),

    # —— 出库 ——
    C("第二批", "D20", "出库管理", "出库管理", "出库列表筛选", "按时间范围筛选；列含消费使用单位、报关单金额、流向企业", "GET /api/outbound", "P0"),
    C("第二批", "D20", "出库管理", "出库管理", "新建出库单", "报关单号、核注清单、流向企业、消费使用单位、出库时间、重量、报关单金额、备注选填", "POST /api/outbound", "P0"),
    C("第二批", "D20", "出库管理", "出库管理", "OCR 出库报关单", "上传出库报关单回填报关单号/核注清单号/流向企业/消费使用单位/重量/金额", "POST /api/ocr/recognize", "P0"),
    C("第二批", "D20", "出库管理", "出库管理", "OCR 字段修正留痕", "全部识别字段可手工修正，写入操作日志（改前改后）", "GET /api/audit-logs", "P0"),
    C("第二批", "D20", "出库管理", "出库管理", "出库审核", "订单审核通过后方可执行", "POST /api/outbound", "P0"),
    C("第二批", "D20", "出库管理", "出库管理", "出库完成", "已审核单据可单次完成并更新出库/完成时间", "POST /api/outbound", "P0"),
    C("第二批", "D20", "出库管理", "出库管理", "出库扣减库存", "按报关单号或生产批次号扣减；余额由已审核流水汇总", "POST /api/outbound", "P0"),

    # —— 生产 ——
    C("第三批", "D26", "生产完工", "生产/完工", "投料/加工合并列表", "投料出库与加工过程同一列表；可按委托方、时间范围筛选", "GET /api/production", "P0"),
    C("第三批", "D26", "生产完工", "生产/完工", "投料出库 1:N", "按生产批次号投料；同一生产批次号可勾选多个报关单号矿物；可用量校验", "POST /api/production", "P0"),
    C("第三批", "D26", "生产完工", "生产/完工", "投料记录出库报关单号", "列表含出库报关单号；矿物按达标矿/报备矿展示", "POST /api/production", "P0"),
    C("第三批", "D26", "生产完工", "生产/完工", "投料筛选与详情", "弹窗可按报关单号、物料、来源国家、堆位筛选；详情展示物料、船名/航次、提单号、来源国家", "GET /api/production/{prodBatchNo}", "P0"),
    C("第三批", "D26", "生产完工", "生产/完工", "加工过程登记", "工序、时间、操作人、异常说明；展示该生产批次下全部报关单号", "POST /api/production", "P0"),
    C("第三批", "D26", "生产完工", "生产/完工", "完工入库仅待检区", "目标堆位只允许待检区；成品库存批次即该生产批次号", "POST /api/production", "P0"),
    C("第三批", "D26", "生产完工", "生产/完工", "谱系追溯", "成品可追溯至各报关单号矿物（1:N）", "GET /api/production/{prodBatchNo}", "P0"),
    C("第三批", "D26", "生产完工", "生产/完工", "消耗产出台账穿透", "原料消耗/成品产出自动汇总，支持穿透到单据", "GET /api/production", "P1"),

    # —— 关务 ——
    C("第三批", "D26", "关务保税", "关务保税", "关务专用操作界面", "关务员权限隔离入口；海关账号在审计中单独标识", "GET /api/customs/reconcile", "P0"),
    C("第三批", "D26", "关务保税", "关务保税", "账实自动对账", "账册库存 vs WMS 账面自动对账（盘点不参与调账）", "GET /api/customs/reconcile", "P0", DONE, DONE, "待确认：对账口径（本系统内比对是否会产生差异）"),
    C("第三批", "D26", "关务保税", "关务保税", "差异清单穿透", "从差异汇总穿透至核注清单、报关单号、生产批次号、入出库单、生产记录", "GET /api/customs/reconcile/{id}/details", "P0"),
    C("第三批", "D26", "关务保税", "关务保税", "对账报表导出", "差异报表导出 Excel", "GET /api/customs/reconcile", "P1"),
    C("第三批", "D26", "关务保税", "关务保税", "联网申报预留", "受控鉴权、回执归档；对端未就绪时交付样例报文与待办", "POST /api/customs/declarations", "P1"),
    C("第三批", "D26", "关务保税", "关务保税", "电子账册核销预留", "电子账册核销端口预留", "POST /api/customs/bonded-books/write-off", "P1"),

    # —— OCR ——
    C("第四批", "D30", "OCR识别", "OCR识别", "OCR 任务中心", "报关单（入/出）/提单/品质单/重量单任务列表与状态", "POST /api/ocr/recognize", "P0"),
    C("第四批", "D30", "OCR识别", "OCR识别", "报关单识别", "入库/出库报关单结构化字段回填", "POST /api/ocr/recognize", "P0"),
    C("第四批", "D30", "OCR识别", "OCR识别", "提单识别", "船名、提单号等回填", "POST /api/ocr/recognize", "P0"),
    C("第四批", "D30", "OCR识别", "OCR识别", "品质单识别", "八项品质与证书号回填", "POST /api/ocr/recognize", "P0"),
    C("第四批", "D30", "OCR识别", "OCR识别", "重量单识别", "湿重/干重/水分回填", "POST /api/ocr/recognize", "P0"),
    C("第四批", "D30", "OCR识别", "OCR识别", "人工修正与归档", "全部识别字段可修正；原图/识别值/修正痕迹可查；附件归档", "POST /api/ocr/recognize", "P0"),
    C("第四批", "D30", "OCR识别", "OCR识别", "OCR 信创 POC 冻结", "UOS + 目标 CPU/JDK 验证识别组件；无包则换厂商或远程识别（过安全评估）", "—", "P0", NA, NA, "D26 前冻结厂商/部署方式", fe=NA, integ=NA),

    # —— 对外对接 ——
    C("第四批", "D30", "对外对接", "对外对接", "对接配置页", "端口清单、启用开关、字段映射入口", "/api/integrations/*", "P0"),
    C("第四批", "D30", "对外对接", "对外对接", "海关端口（必预留）", "单一窗口/联网申报、回执与核销状态；样例报文可调通", "/api/integrations/customs/*", "P0"),
    C("第四批", "D30", "对外对接", "对外对接", "物流端口（必预留）", "提单、船名、到货等协同", "/api/integrations/logistics/*", "P0"),
    C("第四批", "D30", "对外对接", "对外对接", "园区监管端口（扩展）", "监管数据上报预留", "/api/integrations/park/*", "P2"),
    C("第四批", "D30", "对外对接", "对外对接", "企业 ERP 端口（扩展）", "主数据/单据同步预留", "/api/integrations/erp/*", "P2"),
    C("第四批", "D30", "对外对接", "对外对接", "生产系统端口（扩展）", "投料/完工数据协同预留", "/api/integrations/mes/*", "P2"),
    C("第四批", "D30", "对外对接", "对外对接", "接口治理", "白名单、HTTPS/专线、令牌/证书、幂等键、重试队列、调用日志", "/api/integrations/*", "P0", NA, NA, fe=NA),

    # —— 操作日志 ——
    C("第一批", "D10", "操作日志", "操作日志", "日志列表与筛选", "时间段、模块快捷标签、操作人、业务单号、报关单号、生产批次", "GET /api/audit-logs", "P0"),
    C("第一批", "D10", "操作日志", "操作日志", "审计详情", "操作人、时间、对象、内容、改前改后、来源地址、业务单号；关务账号标识", "GET /api/audit-logs/{id}", "P0"),
    C("第一批", "D10", "操作日志", "操作日志", "关键操作全覆盖", "片区变更、OCR 修正、移库、入库、加工登记、权限变更、账号操作均留痕", "GET /api/audit-logs", "P0"),

    # —— 驾驶舱 ——
    C("第四批", "D30", "可视化驾驶舱", "可视化驾驶舱", "科技风 HUD 布局", "深色驾驶舱；KPI 六卡：原料湿/干吨、成品干/湿吨、今日入库、今日出库、待检、预警", "GET /api/dashboard/cockpit", "P0"),
    C("第四批", "D30", "可视化驾驶舱", "可视化驾驶舱", "整园平面图", "默认整园拼接 5 仓；堆位格子叠加在库货物与入库预约；填色与平面图设置一致", "GET /api/dashboard/cockpit", "P0"),
    C("第四批", "D30", "可视化驾驶舱", "可视化驾驶舱", "平面图加载态", "远程数据返回前工业风扫描，至少 2 秒，不先闪默认图", "GET /api/floor-plans", "P1"),
    C("第四批", "D30", "可视化驾驶舱", "可视化驾驶舱", "卫星图开关", "航拍屋面/堆场风格与平面填色切换", "—", "P1"),
    C("第四批", "D30", "可视化驾驶舱", "可视化驾驶舱", "仓/堆位下钻", "点击仓库或堆位镜头放大仅绘该仓；左上角回退带动画；放大后点堆位进管理、点矿物进库存", "GET /api/dashboard/cockpit", "P0"),
    C("第四批", "D30", "可视化驾驶舱", "可视化驾驶舱", "矿物货主分色", "南国铜业绿、五矿有色黄、广西金川蓝；达标浅、报备深、混成品灰；预约格青色描边", "GET /api/dashboard/cockpit", "P0"),
    C("第四批", "D30", "可视化驾驶舱", "可视化驾驶舱", "入出库趋势", "近 7 日入库湿吨 / 出库湿吨", "GET /api/dashboard/cockpit", "P0"),
    C("第四批", "D30", "可视化驾驶舱", "可视化驾驶舱", "作业链路滚动", "加高日志滚动；时间 yyyy-mm-dd HH时；悬停暂停", "GET /api/dashboard/cockpit", "P0"),
    C("第四批", "D30", "可视化驾驶舱", "可视化驾驶舱", "生产数据卡片", "批次号、总量、日期/预计完工、累计重量；点击下钻生产页", "GET /api/dashboard/cockpit", "P1"),
    C("第四批", "D30", "可视化驾驶舱", "可视化驾驶舱", "仓库使用分布", "各仓占用进度条，无饼图", "GET /api/dashboard/cockpit", "P1"),
    C("第四批", "D30", "可视化驾驶舱", "可视化驾驶舱", "预警滚动与下钻", "底栏预警无缝循环、悬停暂停；KPI/预警/指标可下钻至报关单号/生产批次号/单据/日志", "GET /api/dashboard/cockpit", "P0"),

    # —— 第五批交付能力（非业务功能，但需跟踪） ——
    C("第五批", "D45", "部署交付", "部署", "信创环境全量部署", "Nginx + Java + Vastbase + systemd；网络/账号/日志/监控/备份", "—", "P0", NA, NA, fe=NA, integ=NA),
    C("第五批", "D45", "部署交付", "对接联调", "海关/物流真联调或待办清单", "条件未就绪则提交配置说明、样例报文、待办清单", "/api/integrations/customs|logistics", "P0", NA, NA, fe=NA),
    C("第五批", "D45", "部署交付", "验收", "上线初验", "核心链路无阻断、库存计算正确、权限审计有效、备份任务运行", "—", "P0", NA, NA, fe=NA, be=NA, integ=NA),
    C("第五批", "D52", "部署交付", "培训", "分岗位培训考核", "管理员/仓管/关务/运维实操 ≥80 分", "—", "P1", NA, NA, fe=NA, be=NA, integ=NA),
    C("第五批", "D52", "部署交付", "运维", "备份恢复演练", "记录 RPO/RTO，纳入交付", "—", "P0", NA, NA, fe=NA, integ=NA),
    C("第五批", "D55", "部署交付", "验收", "竣工终验与资料移交", "部署手册、运维手册、接口文档、权限矩阵、等保材料齐套", "—", "P0", NA, NA, fe=NA, be=NA, integ=NA),
]


def stage_formula(r: int) -> str:
    """Pipeline: 阻塞 → 已完成 → 测试中 → 对接中 → 开发中 → 待开发 → UI/原型 → 未开始.

    测试中 / 对接中 要求至少有一条开发轨实际「已完成」，避免「全是不适用 + 测试未开始」被当成测试中。
    """
    L, M, N, O = f"L{r}", f"M{r}", f"N{r}", f"O{r}"
    J, K = f"J{r}", f"K{r}"
    done = lambda x: f'OR({x}="已完成",{x}="不适用")'
    return (
        f'IF(COUNTIF(J{r}:O{r},"阻塞")>0,"阻塞",'
        f'IF(COUNTIFS(J{r}:O{r},"<>已完成",J{r}:O{r},"<>不适用")=0,"已完成",'
        f'IF(OR({O}="进行中",AND({done(L)},{done(M)},{done(N)},AND({O}<>"已完成",{O}<>"不适用"),OR({L}="已完成",{M}="已完成",{N}="已完成"))),"测试中",'
        f'IF(OR({N}="进行中",AND({done(L)},{done(M)},AND({N}<>"已完成",{N}<>"不适用"),OR({L}="已完成",{M}="已完成"))),"对接中",'
        f'IF(OR({L}="进行中",{L}="已完成",{M}="进行中",{M}="已完成"),"开发中",'
        f'IF(AND({done(J)},{done(K)}),"待开发",'
        f'IF({K}="进行中","UI设计中",'
        f'IF(OR({J}="进行中",{J}="已完成"),"原型中","未开始"))))))))'
    )


def apply_fill_font(cell, fill=None, font=None, alignment=None, border=True):
    if fill:
        cell.fill = fill
    if font:
        cell.font = font
    cell.alignment = alignment or left
    if border:
        cell.border = thin


def build_guide(wb: Workbook):
    ws = wb.active
    ws.title = "使用说明"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"

    ws.merge_cells("A1:G1")
    ws["A1"].value = "WMS 保税仓储管理系统 · 能力跟踪表"
    ws["A1"].font = font_title
    ws["A1"].fill = fill_title_bar
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 36
    for col in range(2, 8):
        ws.cell(1, col).fill = fill_title_bar

    ws.merge_cells("A2:G2")
    ws["A2"].value = f"对齐项目设计十一大模块 + 公共框架 + 部署交付；生成日期 {date.today().isoformat()}。本表是开发、前后端对接与测试的进度源，填「能力明细」即可，总览自动汇总。"
    ws["A2"].font = font_muted
    ws["A2"].alignment = left
    ws.row_dimensions[2].height = 28

    blocks = [
        ("一、工作表", [
            ("使用说明", "状态口径、阶段算法、填写规则"),
            ("进度总览", "按轨道 / 模块 / 批次自动汇总，含完成率；勿手改带公式的单元格"),
            ("能力明细", "唯一填写入口。每行一条可独立验收的能力"),
            ("页面对照", "原型页面 ↔ 模块 ↔ 能力条数，便于评审走查"),
        ]),
        ("二、六条进度轨道（从左到右）", [
            ("原型设计", "可交互 HTML 原型（wms/dist）是否覆盖该能力。0 期已完成的业务项已预填「已完成」"),
            ("UI设计", "UI设计.md 规范 + 原型视觉是否覆盖。已预填「已完成」"),
            ("前端开发", "正式 React + TypeScript 实现（不含 HTML 原型）"),
            ("后端开发", "正式 Java Spring Boot + Vastbase 实现"),
            ("前后端对接", "联调通过：字段对齐、鉴权、错误码、分页/导出可用"),
            ("测试", "按验收要点完成功能测试（含关键反向用例）并登记结果"),
        ]),
        ("三、状态取值（下拉，勿自造词）", [
            ("未开始", "尚未投入该轨道"),
            ("进行中", "正在做"),
            ("已完成", "该轨道验收通过"),
            ("阻塞", "做不下去，必须填「阻塞原因」"),
            ("不适用", "该能力无此轨道（如纯后端的等保策略无 UI）"),
        ]),
        ("四、「阶段」列（自动，勿手改）", [
            ("未开始", "六轨都未开始"),
            ("原型中 / UI设计中", "仍在设计阶段"),
            ("待开发", "原型+UI 已完成（或不适用），正式开发未开工 ← 当前大部分能力处于此阶段"),
            ("开发中", "前端或后端已开工"),
            ("对接中", "前后端均完成（或无需），正在联调"),
            ("测试中", "联调完成，测试未完成"),
            ("已完成", "六轨均为已完成或不适用"),
            ("阻塞", "任一轨道为阻塞"),
        ]),
        ("五、填写规则", [
            ("谁填", "前端填 前端开发；后端填 后端开发；联调双方共同填 前后端对接；测试填 测试。负责人三列按人填写"),
            ("何时更新", "每条能力状态变化当日更新；每日站会看「进度总览」的阶段分布"),
            ("完成标准", "前端完成 = 页面/交互按 UI 规范可演示（可用 Mock）；后端完成 = 接口可测且写入库规则正确；对接完成 = 去掉 Mock 走真实接口；测试完成 = 验收要点打勾"),
            ("优先级", "P0 主路径/合同验收必须；P1 完整交付；P2 扩展预留（园区/ERP/生产端口）"),
            ("与合同节点", "D10 第一批基础数据权限；D20 仓储主流程；D26 生产关务；D30 测评移交；D45 初验；D55 终验"),
        ]),
        ("六、当前基线（生成时）", [
            ("0 期", "HTML 原型与 UI 设计文档已覆盖业务模块，对应轨道预填「已完成」"),
            ("正式工程", "React / Java 尚未开工，前端、后端、对接、测试预填「未开始」，阶段为「待开发」"),
            ("例外", "工程骨架、等保策略、信创冒烟、接口治理、OCR 信创 POC、部署交付：原型/UI 为「不适用」"),
            ("待确认", "矿源备案「已核销」口径、关务对账口径、平面图现场确认 — 见明细「备注」列，确认后请改备注并继续开发"),
            ("条数", f"当前清单共 {len(CAPS)} 条能力。能力本身增删时才重跑 docs/_gen_capability_tracker.py；日常只改「能力明细」状态，切勿覆盖已填进度"),
        ]),
    ]

    r = 4
    for title, rows in blocks:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
        c = ws.cell(r, 1, title)
        c.font = font_white
        c.fill = fill_primary
        c.alignment = left
        ws.row_dimensions[r].height = 22
        for col in range(2, 8):
            ws.cell(r, col).fill = fill_primary
        r += 1
        ws.cell(r, 1, "项").font = font_bold
        ws.cell(r, 1).fill = fill_sub
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
        ws.cell(r, 2, "说明").font = font_bold
        ws.cell(r, 2).fill = fill_sub
        for col in range(1, 8):
            ws.cell(r, col).fill = fill_sub
            ws.cell(r, col).border = thin
        r += 1
        for name, desc in rows:
            ws.cell(r, 1, name).font = font_bold
            ws.cell(r, 1).alignment = left
            ws.cell(r, 1).border = thin
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
            ws.cell(r, 2, desc).font = font_body
            ws.cell(r, 2).alignment = left
            for col in range(1, 8):
                ws.cell(r, col).border = thin
            ws.row_dimensions[r].height = 36
            r += 1
        r += 1

    ws.column_dimensions["A"].width = 22
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 18
    ws.column_dimensions["B"].width = 88
    ws.freeze_panes = "A3"
    ws.sheet_properties.tabColor = PRIMARY


def rng(col: str, last: int) -> str:
    return f"能力明细!{col}$2:{col}${last}"


def build_overview(wb: Workbook, last: int):
    ws = wb.create_sheet("进度总览")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = ACCENT

    ws.merge_cells("A1:L1")
    ws["A1"].value = "进度总览（公式驱动，请勿在色块内手填）"
    ws["A1"].font = font_title
    ws["A1"].fill = fill_title_bar
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 36
    for col in range(2, 13):
        ws.cell(1, col).fill = fill_title_bar

    ws.merge_cells("A2:L2")
    ws["A2"].value = "完成率分母不含「不适用」。阶段分布来自「能力明细」的自动「阶段」列。"
    ws["A2"].font = font_muted
    ws.row_dimensions[2].height = 20

    kpi_starts = [1, 3, 5, 7, 9, 11]
    labels = ["原型设计", "UI设计", "前端开发", "后端开发", "前后端对接", "测试"]
    cols = ["J", "K", "L", "M", "N", "O"]
    for label, col, c1 in zip(labels, cols, kpi_starts):
        cell = ws.cell(4, c1, label)
        cell.font = font_white
        cell.fill = fill_primary
        cell.alignment = center
        ws.merge_cells(start_row=4, start_column=c1, end_row=4, end_column=c1 + 1)
        ws.cell(4, c1 + 1).fill = fill_primary
        r = rng(col, last)
        rate_cell = ws.cell(5, c1)
        rate_cell.value = (
            f'=IF((COUNTA({r})-COUNTIF({r},"不适用"))=0,"—",'
            f'COUNTIF({r},"已完成")/(COUNTA({r})-COUNTIF({r},"不适用")))'
        )
        rate_cell.number_format = "0%"
        rate_cell.font = font_kpi
        rate_cell.alignment = center
        rate_cell.fill = fill_white
        ws.merge_cells(start_row=5, start_column=c1, end_row=5, end_column=c1 + 1)
        ws.cell(5, c1 + 1).fill = fill_white
        done_cell = ws.cell(6, c1)
        done_cell.value = f'=COUNTIF({r},"已完成")&" 已完成 / "&(COUNTA({r})-COUNTIF({r},"不适用"))&" 适用"'
        done_cell.font = font_muted
        done_cell.alignment = center
        ws.merge_cells(start_row=6, start_column=c1, end_row=6, end_column=c1 + 1)
        for rr in (4, 5, 6):
            for cc in (c1, c1 + 1):
                ws.cell(rr, cc).border = thin
        ws.row_dimensions[4].height = 22
        ws.row_dimensions[5].height = 36
        ws.row_dimensions[6].height = 22

    # Stage distribution
    ws.merge_cells("A8:D8")
    ws["A8"].value = "阶段分布"
    ws["A8"].font = font_white
    ws["A8"].fill = fill_primary
    for col in range(1, 5):
        ws.cell(8, col).fill = fill_primary
        ws.cell(8, col).font = font_white
        ws.cell(8, col).border = thin

    stage_headers = ["阶段", "条数", "占比", "说明"]
    for i, h in enumerate(stage_headers, 1):
        c = ws.cell(9, i, h)
        c.font = font_bold
        c.fill = fill_sub
        c.alignment = center
        c.border = thin

    stage_rows = [
        ("待开发", "原型与 UI 已就绪，等待正式开发"),
        ("开发中", "前端或后端已开工"),
        ("对接中", "前后端完成，正在联调"),
        ("测试中", "联调完成，测试进行中"),
        ("已完成", "六轨均已完成或无需该轨"),
        ("阻塞", "存在阻塞，见明细「阻塞原因」"),
        ("未开始", "设计尚未开始"),
        ("原型中", "原型设计进行中"),
        ("UI设计中", "UI 设计进行中"),
        ("不适用", "整行均为不适用（少见）"),
    ]
    for i, (st, desc) in enumerate(stage_rows):
        rr = 10 + i
        ws.cell(rr, 1, st).font = font_body
        ws.cell(rr, 1).alignment = center
        ws.cell(rr, 2, f'=COUNTIF({rng("P", last)},A{rr})').font = font_body
        ws.cell(rr, 2).alignment = center
        ws.cell(rr, 3, f'=IF(COUNTA({rng("P", last)})=0,0,B{rr}/COUNTA({rng("P", last)}))').font = font_body
        ws.cell(rr, 3).number_format = "0.0%"
        ws.cell(rr, 3).alignment = center
        ws.cell(rr, 4, desc).font = font_muted
        for col in range(1, 5):
            ws.cell(rr, col).border = thin
            ws.cell(rr, col).alignment = left if col == 4 else center

    # Module breakdown starting col F row 8
    ws.merge_cells("F8:L8")
    ws["F8"].value = "按模块"
    ws["F8"].font = font_white
    ws["F8"].fill = fill_primary
    for col in range(6, 13):
        ws.cell(8, col).fill = fill_primary
        ws.cell(8, col).border = thin

    mod_headers = ["模块", "能力数", "待开发", "开发中", "对接中", "测试中", "已完成"]
    for i, h in enumerate(mod_headers, 6):
        c = ws.cell(9, i, h)
        c.font = font_bold
        c.fill = fill_sub
        c.alignment = center
        c.border = thin

    modules = []
    seen = set()
    for cap in CAPS:
        m = cap[2]
        if m not in seen:
            seen.add(m)
            modules.append(m)

    for i, m in enumerate(modules):
        rr = 10 + i
        ws.cell(rr, 6, m).font = font_body
        ws.cell(rr, 7, f'=COUNTIF({rng("D", last)},F{rr})')
        ws.cell(rr, 8, f'=COUNTIFS({rng("D", last)},F{rr},{rng("P", last)},"待开发")')
        ws.cell(rr, 9, f'=COUNTIFS({rng("D", last)},F{rr},{rng("P", last)},"开发中")')
        ws.cell(rr, 10, f'=COUNTIFS({rng("D", last)},F{rr},{rng("P", last)},"对接中")')
        ws.cell(rr, 11, f'=COUNTIFS({rng("D", last)},F{rr},{rng("P", last)},"测试中")')
        ws.cell(rr, 12, f'=COUNTIFS({rng("D", last)},F{rr},{rng("P", last)},"已完成")')
        for col in range(6, 13):
            ws.cell(rr, col).border = thin
            ws.cell(rr, col).alignment = center if col > 6 else left
            ws.cell(rr, col).font = font_body

    last_mod = 10 + len(modules) - 1
    tot = last_mod + 1
    ws.cell(tot, 6, "合计").font = font_bold
    for col in range(7, 13):
        letter = get_column_letter(col)
        ws.cell(tot, col, f"=SUM({letter}10:{letter}{last_mod})").font = font_bold
        ws.cell(tot, col).alignment = center
    for col in range(6, 13):
        ws.cell(tot, col).fill = fill_sub
        ws.cell(tot, col).border = thin
        ws.cell(tot, col).font = font_bold

    note_row = tot + 2
    ws.merge_cells(start_row=note_row, start_column=6, end_row=note_row, end_column=12)
    ws.cell(note_row, 6, "提示：更新「能力明细」中 J–O 列状态后，本页数字与图表自动变化。P0 未完成不得宣称该批次可验收。").font = font_muted

    batch_start = tot + 4
    ws.merge_cells(f"A{batch_start}:H{batch_start}")
    ws.cell(batch_start, 1, "按交付批次").font = font_white
    ws.cell(batch_start, 1).fill = fill_primary
    for col in range(1, 9):
        ws.cell(batch_start, col).fill = fill_primary
        ws.cell(batch_start, col).border = thin
        ws.cell(batch_start, col).font = font_white
    batch_headers2 = ["批次", "目标节点", "能力数", "待开发", "开发中", "对接中", "测试中", "已完成"]
    for i, h in enumerate(batch_headers2, 1):
        c = ws.cell(batch_start + 1, i, h)
        c.font = font_bold
        c.fill = fill_sub
        c.alignment = center
        c.border = thin

    batches = [
        ("启动", "D5"),
        ("第一批", "D10"),
        ("第二批", "D20"),
        ("第三批", "D26"),
        ("第四批", "D30"),
        ("第五批", "D45–D55"),
    ]
    for i, (b, node) in enumerate(batches):
        rr = batch_start + 2 + i
        ws.cell(rr, 1, b).alignment = center
        ws.cell(rr, 2, node).alignment = center
        ws.cell(rr, 3, f'=COUNTIF({rng("B", last)},A{rr})')
        ws.cell(rr, 4, f'=COUNTIFS({rng("B", last)},A{rr},{rng("P", last)},"待开发")')
        ws.cell(rr, 5, f'=COUNTIFS({rng("B", last)},A{rr},{rng("P", last)},"开发中")')
        ws.cell(rr, 6, f'=COUNTIFS({rng("B", last)},A{rr},{rng("P", last)},"对接中")')
        ws.cell(rr, 7, f'=COUNTIFS({rng("B", last)},A{rr},{rng("P", last)},"测试中")')
        ws.cell(rr, 8, f'=COUNTIFS({rng("B", last)},A{rr},{rng("P", last)},"已完成")')
        for col in range(1, 9):
            ws.cell(rr, col).border = thin
            ws.cell(rr, col).font = font_body
            ws.cell(rr, col).alignment = center

    # Chart data (module 待开发 vs 已完成) — use module table
    chart = BarChart()
    chart.type = "col"
    chart.grouping = "stacked"
    chart.title = "各模块阶段（待开发 / 开发中 / 对接中 / 测试中 / 已完成）"
    chart.y_axis.title = "能力条数"
    chart.style = 10
    data = Reference(ws, min_col=8, min_row=9, max_col=12, max_row=last_mod)
    cats = Reference(ws, min_col=6, min_row=10, max_row=last_mod)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.shape = 4
    chart.legend.position = "b"
    chart.width = 22
    chart.height = 10
    ws.add_chart(chart, f"A{batch_start + 10}")

    widths = {
        "A": 16, "B": 14, "C": 12, "D": 36, "E": 14, "F": 18,
        "G": 10, "H": 10, "I": 10, "J": 10, "K": 10, "L": 10, "M": 14,
    }
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    ws.freeze_panes = "A3"
    ws.row_dimensions[8].height = 22


def build_detail(wb: Workbook):
    ws = wb.create_sheet("能力明细", 2)
    ws.sheet_properties.tabColor = SUCCESS

    for col, h in enumerate(HEADERS, 1):
        cell = ws.cell(1, col, h)
        cell.font = font_h
        cell.fill = fill_head
        cell.alignment = center
        cell.border = thin
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"

    dv_status = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=False)
    dv_status.error = "请从下拉选择状态"
    dv_status.errorTitle = "无效状态"
    dv_status.prompt = "未开始 / 进行中 / 已完成 / 阻塞 / 不适用"
    dv_status.promptTitle = "进度状态"
    ws.add_data_validation(dv_status)

    dv_pri = DataValidation(type="list", formula1='"P0,P1,P2"', allow_blank=False)
    ws.add_data_validation(dv_pri)

    dv_batch = DataValidation(type="list", formula1='"启动,第一批,第二批,第三批,第四批,第五批"', allow_blank=False)
    ws.add_data_validation(dv_batch)

    last = len(CAPS) + 1
    dv_status.add(f"J2:O{last}")
    dv_pri.add(f"I2:I{last}")
    dv_batch.add(f"B2:B{last}")

    batch_fills = {
        "启动": PatternFill("solid", fgColor="E8EEF6"),
        "第一批": PatternFill("solid", fgColor="E3F2E9"),
        "第二批": PatternFill("solid", fgColor="E7F1FB"),
        "第三批": PatternFill("solid", fgColor="FFF4D6"),
        "第四批": PatternFill("solid", fgColor="F3E8F7"),
        "第五批": PatternFill("solid", fgColor="F3F5F8"),
    }
    pri_fills = {"P0": fill_p0, "P1": fill_p1, "P2": fill_p2}

    for i, cap in enumerate(CAPS, 1):
        r = i + 1
        batch, node, module, page, name, desc, api, pri, proto, ui, fe, be, integ, test, note = cap
        cid = f"CAP-{i:03d}"
        values = [
            cid, batch, node, module, page, name, desc, api, pri,
            proto, ui, fe, be, integ, test,
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(r, col, val)
            cell.font = font_body
            cell.border = thin
            cell.alignment = center if col in (1, 2, 3, 8, 9) or 10 <= col <= 15 else left
        # stage formula
        sc = ws.cell(r, 16, f"={stage_formula(r)}")
        sc.font = font_bold
        sc.alignment = center
        sc.border = thin
        for col in range(17, 24):
            cell = ws.cell(r, col, note if col == 23 else "")
            cell.font = font_body
            cell.border = thin
            cell.alignment = left
        ws.row_dimensions[r].height = 36
        # batch tint on col B
        ws.cell(r, 2).fill = batch_fills.get(batch, fill_white)
        ws.cell(r, 9).fill = pri_fills.get(pri, fill_white)

    # Conditional formatting for J-O and stage P
    status_cf = [
        ('"已完成"', fill_done, Font(name="微软雅黑", size=10, color=SUCCESS, bold=True)),
        ('"进行中"', fill_doing, Font(name="微软雅黑", size=10, color=WARNING, bold=True)),
        ('"未开始"', fill_todo, Font(name="微软雅黑", size=10, color=MUTED)),
        ('"阻塞"', fill_block, Font(name="微软雅黑", size=10, color=DANGER, bold=True)),
        ('"不适用"', fill_na, Font(name="微软雅黑", size=10, color=ACCENT)),
    ]
    for val, fill, font in status_cf:
        ws.conditional_formatting.add(
            f"J2:O{last}",
            FormulaRule(formula=[f"J2={val}"], fill=fill, font=font),
        )

    stage_cf = [
        ("待开发", fill_wait, Font(name="微软雅黑", size=10, color=PRIMARY, bold=True)),
        ("开发中", fill_doing, Font(name="微软雅黑", size=10, color=WARNING, bold=True)),
        ("对接中", PatternFill("solid", fgColor="F3E8F7"), Font(name="微软雅黑", size=10, color="6B3FA0", bold=True)),
        ("测试中", PatternFill("solid", fgColor="D6F0FF"), Font(name="微软雅黑", size=10, color=ACCENT, bold=True)),
        ("已完成", fill_done, Font(name="微软雅黑", size=10, color=SUCCESS, bold=True)),
        ("阻塞", fill_block, Font(name="微软雅黑", size=10, color=DANGER, bold=True)),
        ("未开始", fill_todo, Font(name="微软雅黑", size=10, color=MUTED)),
        ("原型中", fill_doing, Font(name="微软雅黑", size=10, color=WARNING)),
        ("UI设计中", fill_doing, Font(name="微软雅黑", size=10, color=WARNING)),
        ("不适用", fill_na, Font(name="微软雅黑", size=10, color=ACCENT)),
    ]
    for val, fill, font in stage_cf:
        ws.conditional_formatting.add(
            f"P2:P{last}",
            FormulaRule(formula=[f'P2="{val}"'], fill=fill, font=font),
        )

    tab = Table(displayName="Caps", ref=f"A1:W{last}")
    tab.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False,
        showLastColumn=False, showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(tab)

    widths = {
        "A": 11, "B": 10, "C": 10, "D": 16, "E": 22, "F": 24, "G": 52,
        "H": 42, "I": 8, "J": 12, "K": 12, "L": 12, "M": 12, "N": 12, "O": 10,
        "P": 12, "Q": 12, "R": 12, "S": 12, "T": 12, "U": 12, "V": 22, "W": 40,
    }
    for k, v in widths.items():
        ws.column_dimensions[k].width = v

    ws.sheet_view.zoomScale = 90
    ws.print_title_rows = "1:1"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.oddHeader.left.text = "WMS 能力明细"
    ws.oddFooter.right.text = "第 &P 页 / 共 &N 页"

    # comments on header
    ws["J1"].comment = None
    return last


def build_pages(wb: Workbook, last: int):
    ws = wb.create_sheet("页面对照")
    ws.sheet_properties.tabColor = WARNING
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:F1")
    ws["A1"].value = "原型页面对照（便于评审走查与前后端分文件）"
    ws["A1"].font = font_title
    ws["A1"].fill = fill_title_bar
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 36
    for col in range(2, 7):
        ws.cell(1, col).fill = fill_title_bar

    headers = ["导航分组", "页面 ID", "页面名称", "对应模块", "能力条数", "主要接口"]
    for i, h in enumerate(headers, 1):
        c = ws.cell(3, i, h)
        c.font = font_h
        c.fill = fill_head
        c.alignment = center
        c.border = thin

    pages = [
        ("概览", "dashboard", "可视化驾驶舱", "可视化驾驶舱", "GET /api/dashboard/cockpit；GET /api/floor-plans"),
        ("基础数据", "warehouses", "仓库管理", "仓库管理", "GET/POST /api/warehouses"),
        ("基础数据", "stacks", "堆位管理", "堆位管理", "GET/POST /api/warehouses/stacks；片区 PUT .../area-label"),
        ("基础数据", "materials", "物料档案 / 矿源备案", "物料档案、矿源备案", "GET/POST /api/materials；/api/material-sources/filings"),
        ("基础数据", "partners", "往来主体", "往来主体", "GET/POST /api/partners"),
        ("基础数据", "bonded", "保税账册", "保税账册", "GET/POST /api/bonded-books"),
        ("基础数据", "settings", "系统设置（账号/角色/参数/平面图）", "系统设置、平面图设置、公共框架", "GET/POST /api/auth/users；roles；floor-plans"),
        ("仓储作业", "inbound", "入库管理", "入库管理", "GET/POST /api/inbound；POST /api/ocr/recognize"),
        ("仓储作业", "transfer", "库内移库", "库内移库", "GET/POST /api/transfers；POST /api/transfers/batch"),
        ("仓储作业", "inventory", "库存查询", "库存查询", "GET /api/inventory；batches/{batchNo}；ledger"),
        ("仓储作业", "stocktake", "盘点管理", "盘点管理", "GET/POST /api/stocktakes；onhand-export"),
        ("仓储作业", "alert", "库存预警", "库存预警", "GET /api/inventory/alerts"),
        ("仓储作业", "outbound", "出库管理", "出库管理", "GET/POST /api/outbound"),
        ("生产与关务", "production", "生产/完工", "生产完工", "GET/POST /api/production"),
        ("生产与关务", "customs", "关务保税", "关务保税", "GET /api/customs/reconcile"),
        ("生产与关务", "ocr", "OCR识别", "OCR识别", "POST /api/ocr/recognize"),
        ("系统", "integration", "对外对接", "对外对接", "/api/integrations/*"),
        ("系统", "audit", "操作日志", "操作日志", "GET /api/audit-logs"),
        ("—", "login", "登录页", "公共框架", "POST /api/auth/login"),
        ("—", "—", "工程 / 信创 / 部署（无独立页面）", "公共框架、部署交付", "健康检查 / 部署脚本"),
    ]

    for i, (grp, pid, pname, mods, apis) in enumerate(pages):
        r = 4 + i
        ws.cell(r, 1, grp)
        ws.cell(r, 2, pid)
        ws.cell(r, 3, pname)
        ws.cell(r, 4, mods)
        ws.cell(r, 6, apis)
        for col in range(1, 7):
            ws.cell(r, col).font = font_body
            ws.cell(r, col).border = thin
            ws.cell(r, col).alignment = left if col in (3, 4, 6) else center
        ws.row_dimensions[r].height = 28
    d = rng("D", last)
    e = rng("E", last)
    formulas = {
        4: f'=COUNTIF({d},"可视化驾驶舱")',
        5: f'=COUNTIF({d},"仓库管理")',
        6: f'=COUNTIF({d},"堆位管理")',
        7: f'=COUNTIF({d},"物料档案")+COUNTIF({d},"矿源备案")',
        8: f'=COUNTIF({d},"往来主体")',
        9: f'=COUNTIF({d},"保税账册")',
        10: f'=COUNTIF({d},"系统设置")+COUNTIF({d},"平面图设置")',
        11: f'=COUNTIF({d},"入库管理")',
        12: f'=COUNTIF({d},"库内移库")',
        13: f'=COUNTIF({d},"库存查询")',
        14: f'=COUNTIF({d},"盘点管理")',
        15: f'=COUNTIF({d},"库存预警")',
        16: f'=COUNTIF({d},"出库管理")',
        17: f'=COUNTIF({d},"生产完工")',
        18: f'=COUNTIF({d},"关务保税")',
        19: f'=COUNTIF({d},"OCR识别")',
        20: f'=COUNTIF({d},"对外对接")',
        21: f'=COUNTIF({d},"操作日志")',
        22: f'=COUNTIFS({d},"公共框架",{e},"<>工程")',
        23: f'=COUNTIF({d},"部署交付")+COUNTIFS({d},"公共框架",{e},"工程")',
    }
    for r, f in formulas.items():
        ws.cell(r, 5, f)
        ws.cell(r, 5).alignment = center
        ws.cell(r, 5).font = font_bold

    note_r = 25
    ws.merge_cells(start_row=note_r, start_column=1, end_row=note_r, end_column=6)
    ws.cell(note_r, 1, "走查建议：按本表页面顺序打开 wms/dist/index.html，对照「能力明细」中该模块全部 P0 项。关务员角色应能看到全部菜单；仓管员仅仓储作业相关。").font = font_muted

    for col, w in zip("ABCDEF", [14, 16, 36, 32, 12, 62]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"


def main():
    last = len(CAPS) + 1
    wb = Workbook()
    build_guide(wb)
    build_detail(wb)
    build_overview(wb, last)
    build_pages(wb, last)

    # Reorder sheets: 使用说明, 进度总览, 能力明细, 页面对照
    order = ["使用说明", "进度总览", "能力明细", "页面对照"]
    for i, name in enumerate(order):
        wb.move_sheet(name, offset=i - wb.sheetnames.index(name))

    wb.properties.creator = "WMS 能力跟踪"
    wb.properties.title = "WMS 系统能力跟踪表"
    wb.properties.description = "细分系统能力，跟踪原型、UI、前后端开发、对接与测试进度"
    wb.save(OUT)
    print(f"Wrote {OUT}  rows={len(CAPS)}")


if __name__ == "__main__":
    main()
